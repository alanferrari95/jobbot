"""Resolve a (tailored) CV spec against base_cv.yaml and render it to DOCX + PDF."""

from __future__ import annotations

import copy
import html
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"


# ---------------------------------------------------------------- resolve

def all_bullets(base: dict) -> dict[str, str]:
    out = {}
    for role in base["experience"]:
        for g in role["groups"]:
            for b in g["bullets"]:
                out[b["id"]] = b["text"]
    for p in base["projects"]:
        for b in p["bullets"]:
            out[b["id"]] = b["text"]
    return out


def resolve(base: dict, spec: dict | None = None) -> dict:
    """Apply a tailoring spec to the base CV. Unknown bullet ids raise KeyError."""
    spec = spec or {}
    cv = copy.deepcopy(base)
    cv["title"] = spec.get("title", base["title"])
    cv["summary"] = spec.get("summary", base["summary"])

    order = spec.get("experience", {}) or {}
    for role in cv["experience"]:
        if role["id"] not in order:
            continue
        wanted = order[role["id"]] or []
        by_id = {b["id"]: (gi, b) for gi, g in enumerate(role["groups"]) for b in g["bullets"]}
        for bid in wanted:
            if bid not in by_id:
                raise KeyError(f"unknown bullet id {bid!r} for role {role['id']!r}")
        groups: list[dict] = []
        index: dict[int, dict] = {}
        for bid in wanted:
            gi, b = by_id[bid]
            if gi not in index:
                g = {k: v for k, v in role["groups"][gi].items() if k != "bullets"}
                g["bullets"] = []
                index[gi] = g
                groups.append(g)
            index[gi]["bullets"].append(b)
        role["groups"] = groups

    if "project" in spec:
        known = {b["id"] for p in base["projects"] for b in p["bullets"]}
        for bid in spec["project"]:
            if bid not in known:
                raise KeyError(f"unknown project bullet id {bid!r}")
        # Bullet ids keep their project; a project with no listed ids is dropped.
        for p in cv["projects"]:
            by_id = {b["id"]: b for b in p["bullets"]}
            p["bullets"] = [by_id[bid] for bid in spec["project"] if bid in by_id]

    if "skills" in spec:
        cv["skills"] = spec["skills"]
    return cv


# ---------------------------------------------------------------- render

class _Writer:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.7)
        sec.top_margin = sec.bottom_margin = Cm(1.4)
        self.width = sec.page_width - sec.left_margin - sec.right_margin
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.05

    def para(self, text="", size=10, bold=False, color=None, italic=False, before=0, after=0):
        p = self.doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        return p

    def link(self, p, url, text, size=9.5):
        r_id = p.part.relate_to(
            url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "555555")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(color)
        rpr.append(sz)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    def sep(self, p, size=9.5):
        r = p.add_run("  ·  ")
        r.font.size = Pt(size)
        r.font.color.rgb = GREY

    def section(self, title):
        p = self.para(title.upper(), size=11, bold=True, color=ACCENT, before=8, after=3)
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F4E79"}.items():
            b.set(qn(k), v)
        bdr.append(b)
        pPr.append(bdr)

    def role(self, left, right, sub=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(self.width, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(left)
        r.bold = True
        r.font.size = Pt(10.5)
        r2 = p.add_run("\t" + right)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = GREY
        if sub:
            self.para(sub, size=9.5, italic=True, color=GREY, after=1)

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.55)
        p.add_run(text)

    def skills_line(self, label, items):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run(label + ": ")
        r.bold = True
        p.add_run(", ".join(items))


def render_docx(cv: dict, out: Path) -> Path:
    w = _Writer()
    c = cv["contact"]
    w.para(c["name"], size=20, bold=True, color=ACCENT)
    w.para(cv["title"], size=11.5, bold=True, after=2)
    p = w.para()
    links = [(f"mailto:{c['email']}", c["email"])]
    if c.get("phone"):
        links.append(("tel:" + c["phone"].replace(" ", ""), c["phone"]))
    if c.get("linkedin"):
        links.append((c["linkedin"], c["linkedin"].replace("https://www.", "")))
    if c.get("github"):
        links.append((c["github"], c["github"].replace("https://", "")))
    for n, (url, text) in enumerate(links):
        if n:
            w.sep(p)
        w.link(p, url, text)
    w.para(c["location_line"], size=9.5, color=GREY, after=2)

    w.section("Summary")
    w.para(" ".join(cv["summary"].split()), after=2)

    w.section("Experience")
    for role in cv["experience"]:
        if not any(g["bullets"] for g in role["groups"]):  # role dropped by the spec
            continue
        w.role(f"{role['company']}: {role['role']}", f"{role['dates']} · {role['location']}", role.get("blurb"))
        for g in role["groups"]:
            if g.get("client") and g["bullets"]:
                w.para(g["client"], bold=True, before=3, after=1)
            for b in g["bullets"]:
                w.bullet(b["text"])

    projects = [pr for pr in cv["projects"] if pr["bullets"]]
    if projects:
        w.section("Projects")
    for pr in projects:
        w.role(pr["name"], pr["dates"], pr.get("url", "").replace("https://", "") or pr.get("note"))
        for b in pr["bullets"]:
            w.bullet(b["text"])

    w.section("Skills")
    for s in cv["skills"]:
        w.skills_line(s["label"], s["items"])

    w.section("Certifications")
    for cert in cv["certifications"]:
        w.bullet(cert)

    w.section("Education")
    for e in cv["education"]:
        w.role(f"{e['school']}: {e['degree']}", e["dates"])

    w.section("Languages")
    w.para(cv["languages"])

    out.parent.mkdir(parents=True, exist_ok=True)
    w.doc.save(out)
    return out


# ---------------------------------------------------------------- PDF
# JOBBOT_PDF picks the converter: word (Windows, COM), libreoffice (any OS, headless), basic (built-in,
# plain layout, used by the demo and CI), or auto (default): Word if installed, else LibreOffice, else basic.

def _word_installed() -> bool:
    if sys.platform != "win32":
        return False
    import winreg
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application"))
        return True
    except OSError:
        return False


def _soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(p).exists():
            return p
    return None


def _pdf_word(docx: Path, pdf: Path) -> Path:
    """Convert with Microsoft Word through COM (Windows)."""
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"try {{ $d = $w.Documents.Open('{docx.resolve()}', $false, $true); "
        f"$d.SaveAs2('{pdf.resolve()}', 17); $d.Close($false) }} finally {{ $w.Quit() }}"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, capture_output=True)
    return pdf


def _pdf_libreoffice(docx: Path, pdf: Path) -> Path:
    soffice = _soffice()
    if not soffice:
        raise RuntimeError("LibreOffice not found: install it or set JOBBOT_PDF=basic")
    with tempfile.TemporaryDirectory() as tmp:  # LibreOffice names the output after the input
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx.resolve())],
                       check=True, capture_output=True, timeout=180)
        shutil.move(str(Path(tmp) / f"{docx.stem}.pdf"), pdf)
    return pdf


def _pdf_basic(cv: dict, pdf: Path) -> Path:
    """Plain, ATS-readable PDF drawn with PyMuPDF from the same resolved CV (no Office needed)."""
    import pymupdf

    e = html.escape
    c = cv["contact"]
    links = " · ".join(e(c[k]) for k in ("email", "phone", "linkedin", "github") if c.get(k))
    parts = [f"<h1>{e(c['name'])}</h1><p><b>{e(cv['title'])}</b></p><p class=g>{links}</p>",
             f"<p class=g>{e(c.get('location_line', ''))}</p>",
             f"<h2>SUMMARY</h2><p>{e(' '.join(cv['summary'].split()))}</p>", "<h2>EXPERIENCE</h2>"]
    for role in cv["experience"]:
        if not any(g["bullets"] for g in role["groups"]):  # role dropped by the spec
            continue
        parts.append(f"<h3>{e(role['company'])}: {e(role['role'])} <span class=g>· {e(role['dates'])} · "
                     f"{e(role['location'])}</span></h3>")
        for g in role["groups"]:
            if g.get("client") and g["bullets"]:
                parts.append(f"<p><b>{e(g['client'])}</b></p>")
            parts.append("<ul>" + "".join(f"<li>{e(b['text'])}</li>" for b in g["bullets"]) + "</ul>")
    projects = [pr for pr in cv["projects"] if pr["bullets"]]
    if projects:
        parts.append("<h2>PROJECTS</h2>")
    for pr in projects:
        parts.append(f"<h3>{e(pr['name'])} <span class=g>· {e(pr['dates'])}</span></h3>"
                     + "<ul>" + "".join(f"<li>{e(b['text'])}</li>" for b in pr["bullets"]) + "</ul>")
    parts.append("<h2>SKILLS</h2>" + "".join(f"<p><b>{e(s['label'])}:</b> {e(', '.join(s['items']))}</p>"
                                             for s in cv["skills"]))
    parts.append("<h2>CERTIFICATIONS</h2><ul>" + "".join(f"<li>{e(x)}</li>" for x in cv["certifications"]) + "</ul>")
    parts.append("<h2>EDUCATION</h2>" + "".join(f"<p>{e(x['school'])}: {e(x['degree'])} · {e(x['dates'])}</p>"
                                                 for x in cv["education"]))
    parts.append(f"<h2>LANGUAGES</h2><p>{e(cv['languages'])}</p>")
    css = ("* {font-family: sans-serif; font-size: 9.5pt;} h1 {font-size: 18pt; color: #1F4E79;} "
           "h2 {font-size: 11pt; color: #1F4E79; margin-top: 8pt;} h3 {font-size: 10pt; margin-top: 5pt;} "
           ".g {color: #555555;} ul {margin-top: 1pt;} p {margin: 1pt 0;}")
    story = pymupdf.Story(html="".join(parts), user_css=css)
    pdf.parent.mkdir(parents=True, exist_ok=True)
    writer = pymupdf.DocumentWriter(str(pdf))
    page = pymupdf.paper_rect("a4")
    more = True
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(page + (48, 40, -48, -40))
        story.draw(dev)
        writer.end_page()
    writer.close()
    return pdf


def pdf_engine() -> str:
    engine = os.environ.get("JOBBOT_PDF", "auto")
    if engine != "auto":
        return engine
    if _word_installed():
        return "word"
    return "libreoffice" if _soffice() else "basic"


def docx_to_pdf(docx: Path, pdf: Path, engine: str | None = None) -> Path:
    engine = engine or pdf_engine()
    if engine == "word":
        return _pdf_word(docx, pdf)
    if engine == "libreoffice":
        return _pdf_libreoffice(docx, pdf)
    raise ValueError(f"unknown PDF engine {engine!r} (word | libreoffice | basic | auto)")


def render(cv: dict, out_dir: Path, stem: str) -> Path:
    docx = render_docx(cv, out_dir / f"{stem}.docx")
    engine = pdf_engine()
    if engine == "basic":
        return _pdf_basic(cv, out_dir / f"{stem}.pdf")
    return docx_to_pdf(docx, out_dir / f"{stem}.pdf", engine)
